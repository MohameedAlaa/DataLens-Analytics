"""Regression tests: a categorical column holding numeric 1 and string "1".

Identical string representations are one displayed categorical label, so
precomputed options, category search and /filter must agree on a single
combined entry and return every record the label represents.
"""

import pandas as pd
import pytest

from app import app
from api_helpers import compute_filter_options, filter_dataframe, search_category_values

DATASET_ID = "mixed-type-dataset"


def _mixed_df() -> pd.DataFrame:
    # object column: numeric 1 (x1), string "1" (x2), "abc", and a null
    return pd.DataFrame(
        {
            "id": [1, 2, 3, 4, 5],
            "code": [1, "1", "1", "abc", None],
            "score": [10, 20, 30, 40, 50],
        }
    )


@pytest.fixture
def mixed_dataset(monkeypatch):
    """Register a dataset whose served DataFrame has a genuinely mixed-type column.

    The disk engine persists via parquet, which cannot store an object column
    holding both int and str, so a placeholder is registered (for metadata and
    readiness) and ``get_dataset`` serves the mixed frame for this id only.
    """
    df = _mixed_df()
    manager = app.state.dataset_manager
    placeholder = pd.DataFrame({"id": [1]})
    manager.register_dataset(df=placeholder, filename="mixed.csv", dataset_id=DATASET_ID)
    manager.update_indexing_status(DATASET_ID, "READY")

    original = manager.get_dataset

    def _get_dataset(dataset_id):
        if dataset_id == DATASET_ID:
            return df.copy()
        return original(dataset_id)

    monkeypatch.setattr(manager, "get_dataset", _get_dataset)
    return df


def test_precomputed_options_have_no_duplicate_labels():
    df = _mixed_df()
    opts = compute_filter_options(df)["code"]["top_values"]
    labels = [o["value"] for o in opts]
    assert len(labels) == len(set(labels)), f"duplicate labels: {labels}"
    counts = {o["value"]: o["count"] for o in opts}
    assert counts == {"1": 3, "abc": 1}


def test_search_helper_returns_single_combined_label():
    df = _mixed_df()
    res = search_category_values(df, "code", "1")
    assert res["values"] == [{"value": "1", "count": 3}]
    assert res["total_matches"] == 1
    # precomputed and search must agree on the count
    pre = {o["value"]: o["count"] for o in compute_filter_options(df)["code"]["top_values"]}
    assert pre["1"] == res["values"][0]["count"]


def test_filter_dataframe_label_matches_numeric_and_string_rows():
    df = _mixed_df()
    snapshot = df.copy(deep=True)
    out = filter_dataframe(df, {"code": ["1"]})
    assert sorted(out["id"].tolist()) == [1, 2, 3]
    # input frame is not mutated (no in-place dtype changes)
    pd.testing.assert_frame_equal(df, snapshot)
    # nulls are never matched
    assert 5 not in out["id"].tolist()


def test_api_search_and_filter_roundtrip(client, mixed_dataset):
    resp = client.get(
        "/api/categories",
        params={"dataset_id": DATASET_ID, "column": "code", "query": "1"},
    )
    assert resp.status_code == 200
    assert resp.json()["values"] == [{"value": "1", "count": 3}]

    filt = client.post(
        "/filter",
        json={"dataset_id": DATASET_ID, "filters": {"code": ["1"]}},
    )
    assert filt.status_code == 200
    body = filt.json()
    assert body["row_count"] == 3
    assert sorted(r["id"] for r in body["filtered_data"]) == [1, 2, 3]

    other = client.post(
        "/filter",
        json={"dataset_id": DATASET_ID, "filters": {"code": ["abc"]}},
    )
    assert other.json()["row_count"] == 1


def test_numeric_filter_semantics_unchanged(client, mixed_dataset):
    df = _mixed_df()
    # list filter on a genuinely numeric column keeps exact numeric matching
    assert filter_dataframe(df, {"score": [10, 30]})["id"].tolist() == [1, 3]
    # numeric range / min / max semantics unchanged
    assert filter_dataframe(df, {"score": {"min": 20, "max": 40}})["id"].tolist() == [2, 3, 4]
    assert filter_dataframe(df, {"score": {"range": [10, 20]}})["id"].tolist() == [1, 2]
    resp = client.post(
        "/filter",
        json={"dataset_id": DATASET_ID, "filters": {"score": {"min": 40}}},
    )
    assert resp.json()["row_count"] == 2


def _make_mixed_excel() -> bytes:
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["id", "code", "score"])
    # Row 1: numeric 1
    ws.append([1, 1, 10])
    # Row 2: string "1"
    ws.cell(row=3, column=1, value=2)
    ws.cell(row=3, column=2, value="1")
    ws.cell(row=3, column=3, value=20)
    # Row 3: string "abc"
    ws.append([3, "abc", 30])
    # Row 4: null
    ws.append([4, None, 40])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_excel_mixed_type_persistence_with_disk_storage(tmp_path):
    import io
    from dataset_manager import DatasetManager, DiskStorageEngine
    excel_bytes = _make_mixed_excel()
    df = pd.read_excel(io.BytesIO(excel_bytes))

    disk_engine = DiskStorageEngine(storage_dir=str(tmp_path))
    manager = DatasetManager(engine=disk_engine)

    # Without the DiskStorageEngine fix, registering df with mixed object types fails
    manager.register_dataset(df=df, filename="mixed.xlsx", dataset_id="disk-mixed")

    loaded = manager.get_dataset("disk-mixed")
    assert loaded is not None
    # Code values converted cleanly to string/None
    codes = loaded["code"].tolist()
    assert codes[0] == "1"
    assert codes[1] == "1"
    assert codes[2] == "abc"
    assert pd.isna(codes[3])


def test_excel_upload_with_disk_storage_end_to_end(client, tmp_path, monkeypatch):
    from dataset_manager import DatasetManager, DiskStorageEngine
    disk_engine = DiskStorageEngine(storage_dir=str(tmp_path))
    disk_manager = DatasetManager(engine=disk_engine)
    monkeypatch.setattr(app.state, "dataset_manager", disk_manager)

    excel_bytes = _make_mixed_excel()
    resp = client.post(
        "/upload",
        files={"file": ("mixed.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    dataset_id = resp.json()["dataset_id"]

    disk_manager.update_indexing_status(dataset_id, "READY")

    # Search finds the combined entry (row 1 numeric 1 and row 2 string "1")
    search_resp = client.get(
        "/api/categories",
        params={"dataset_id": dataset_id, "column": "code", "query": "1"},
    )
    assert search_resp.status_code == 200
    assert search_resp.json()["values"] == [{"value": "1", "count": 2}]

    # Filter matches both numeric and string rows
    filter_resp = client.post(
        "/filter",
        json={"dataset_id": dataset_id, "filters": {"code": ["1"]}},
    )
    assert filter_resp.status_code == 200
    assert filter_resp.json()["row_count"] == 2


def test_unrelated_filesystem_error_does_not_trigger_mixed_type_fallback(tmp_path, monkeypatch):
    """Verify that filesystem / permission errors do NOT invoke the mixed-type normalization fallback."""
    from dataset_manager import DiskStorageEngine
    engine = DiskStorageEngine(storage_dir=str(tmp_path))
    df = pd.DataFrame({"text_col": ["val1", "val2"]})

    fallback_called = False
    original_map = pd.Series.map

    def tracking_map(*args, **kwargs):
        nonlocal fallback_called
        fallback_called = True
        return original_map(*args, **kwargs)

    # Force to_parquet to raise an OSError (e.g. disk full / permission denied)
    def failing_to_parquet(*args, **kwargs):
        raise OSError("Disk quota exceeded")

    monkeypatch.setattr(pd.DataFrame, "to_parquet", failing_to_parquet)
    monkeypatch.setattr(pd.Series, "map", tracking_map)

    with pytest.raises(OSError, match="Disk quota exceeded"):
        engine.store("err-dataset", df, {"filename": "test.csv"})

    assert fallback_called is False, "Mixed-type fallback should not be triggered on OSError"


def test_failed_serialization_leaves_no_orphaned_files(tmp_path, monkeypatch):
    """Verify that when serialization fails completely, no orphaned or 0-byte files remain."""
    from dataset_manager import DiskStorageEngine
    engine = DiskStorageEngine(storage_dir=str(tmp_path))
    df = pd.DataFrame({"a": [1, 2]})

    import pyarrow as pa
    def fatal_arrow_error(self, path, *args, **kwargs):
        # Simulate partial file creation on disk before error
        with open(path, "wb") as f:
            f.write(b"partial bytes")
        raise pa.ArrowInvalid("Fatal serialization error")

    monkeypatch.setattr(pd.DataFrame, "to_parquet", fatal_arrow_error)

    with pytest.raises(pa.ArrowInvalid, match="Fatal serialization error"):
        engine.store("fatal-dataset", df, {"filename": "test.csv"})

    # Check directory contents: no .parquet, .meta.json, or .tmp_* files
    remaining_files = list(tmp_path.iterdir())
    assert remaining_files == [], f"Expected clean storage directory, but found: {remaining_files}"
    assert engine.exists("fatal-dataset") is False
    assert engine.list_ids() == []


def test_failed_overwrite_preserves_existing_valid_dataset(tmp_path, monkeypatch):
    """Verify that a failed replacement write leaves the existing valid dataset completely intact."""
    from dataset_manager import DiskStorageEngine
    engine = DiskStorageEngine(storage_dir=str(tmp_path))
    original_df = pd.DataFrame({"id": [1, 2], "val": ["A", "B"]})
    original_meta = {"dataset_id": "ds_persist", "rows": 2, "version": 1}

    # Store initial valid dataset
    engine.store("ds_persist", original_df, original_meta)
    assert engine.exists("ds_persist") is True
    pd.testing.assert_frame_equal(engine.get("ds_persist"), original_df)
    assert engine.get_metadata("ds_persist")["version"] == 1

    # Now attempt an update that fails during Parquet write
    replacement_df = pd.DataFrame({"id": [9, 10], "val": ["X", "Y"]})
    replacement_meta = {"dataset_id": "ds_persist", "rows": 2, "version": 2}

    import pyarrow as pa
    def failing_to_parquet(self, path, *args, **kwargs):
        # Open file to simulate partial write then fail
        with open(path, "wb") as f:
            f.write(b"corrupt partial parquet")
        raise pa.ArrowInvalid("Corrupted serialization during update")

    monkeypatch.setattr(pd.DataFrame, "to_parquet", failing_to_parquet)

    with pytest.raises(pa.ArrowInvalid):
        engine.store("ds_persist", replacement_df, replacement_meta)

    # Existing dataset must be preserved without corruption or data loss
    assert engine.exists("ds_persist") is True
    pd.testing.assert_frame_equal(engine.get("ds_persist"), original_df)
    assert engine.get_metadata("ds_persist")["version"] == 1



def test_metadata_and_parquet_atomic_consistency(tmp_path):
    """Verify that store() writes both parquet and metadata atomically."""
    from dataset_manager import DiskStorageEngine
    engine = DiskStorageEngine(storage_dir=str(tmp_path))
    df = pd.DataFrame({"id": [1, 2], "score": [10.5, 20.0]})
    meta = {"dataset_id": "atomic_ds", "indexing_status": "READY"}

    engine.store("atomic_ds", df, meta)

    df_path = engine._df_path("atomic_ds")
    meta_path = engine._meta_path("atomic_ds")
    assert df_path.is_file()
    assert meta_path.is_file()
    assert engine.exists("atomic_ds") is True
    assert engine.get("atomic_ds").shape == (2, 2)
    assert engine.get_metadata("atomic_ds")["indexing_status"] == "READY"


def test_second_replace_failure_on_new_dataset_leaves_no_orphaned_parquet(tmp_path, monkeypatch):
    """Verify that if the second os.replace (metadata) fails on a new dataset, no orphaned Parquet remains."""
    import os
    from dataset_manager import DiskStorageEngine
    engine = DiskStorageEngine(storage_dir=str(tmp_path))
    df = pd.DataFrame({"id": [1, 2], "val": ["A", "B"]})
    meta = {"dataset_id": "new_ds", "rows": 2}

    meta_target = str(engine._meta_path("new_ds"))
    original_replace = os.replace

    def fail_on_meta_replace(src, dst, *args, **kwargs):
        if str(dst) == meta_target or str(dst).endswith("new_ds.meta.json"):
            raise OSError("Injected disk failure during metadata os.replace")
        return original_replace(src, dst, *args, **kwargs)

    monkeypatch.setattr(os, "replace", fail_on_meta_replace)

    with pytest.raises(OSError, match="Injected disk failure during metadata os.replace"):
        engine.store("new_ds", df, meta)

    # Verify no orphaned Parquet or metadata file exists
    assert engine.exists("new_ds") is False
    assert engine.list_ids() == []
    remaining = [p.name for p in tmp_path.iterdir()]
    assert remaining == [], f"Expected 0 leftover files, found: {remaining}"


def test_second_replace_failure_on_update_rolls_back_to_original_valid_dataset(tmp_path, monkeypatch):
    """Verify that if the second os.replace fails during an update, the original valid dataset is fully restored."""
    import os
    from dataset_manager import DiskStorageEngine
    engine = DiskStorageEngine(storage_dir=str(tmp_path))
    v1_df = pd.DataFrame({"id": [1, 2], "val": ["V1_A", "V1_B"]})
    v1_meta = {"dataset_id": "update_ds", "version": 1}

    # Store initial valid dataset
    engine.store("update_ds", v1_df, v1_meta)
    assert engine.exists("update_ds") is True
    pd.testing.assert_frame_equal(engine.get("update_ds"), v1_df)

    v2_df = pd.DataFrame({"id": [10, 20], "val": ["V2_X", "V2_Y"]})
    v2_meta = {"dataset_id": "update_ds", "version": 2}

    meta_target = str(engine._meta_path("update_ds"))
    original_replace = os.replace

    def fail_on_meta_replace(src, dst, *args, **kwargs):
        if ".tmp_" in str(src) and (str(dst) == meta_target or str(dst).endswith("update_ds.meta.json")):
            raise OSError("Injected disk failure during metadata os.replace on update")
        return original_replace(src, dst, *args, **kwargs)

    monkeypatch.setattr(os, "replace", fail_on_meta_replace)

    with pytest.raises(OSError, match="Injected disk failure during metadata os.replace on update"):
        engine.store("update_ds", v2_df, v2_meta)

    # Verify full rollback to v1: both Parquet and metadata match v1
    assert engine.exists("update_ds") is True
    pd.testing.assert_frame_equal(engine.get("update_ds"), v1_df)
    assert engine.get_metadata("update_ds")["version"] == 1

    # Verify no leftover tmp or bak files
    temp_or_bak = [p.name for p in tmp_path.iterdir() if ".tmp_" in p.name or ".bak_" in p.name]
    assert temp_or_bak == [], f"Leftover tmp/bak files: {temp_or_bak}"


def test_metadata_replace_fails_and_meta_rollback_fails_parquet_is_still_restored(tmp_path, monkeypatch):
    """Verify that if restoring metadata backup fails during rollback, Parquet rollback still runs independently."""
    import os
    from dataset_manager import DiskStorageEngine
    engine = DiskStorageEngine(storage_dir=str(tmp_path))
    v1_df = pd.DataFrame({"id": [1, 2], "val": ["V1_A", "V1_B"]})
    v1_meta = {"dataset_id": "indep_ds", "version": 1}

    engine.store("indep_ds", v1_df, v1_meta)
    assert engine.exists("indep_ds") is True

    v2_df = pd.DataFrame({"id": [10, 20], "val": ["V2_X", "V2_Y"]})
    v2_meta = {"dataset_id": "indep_ds", "version": 2}

    meta_target = str(engine._meta_path("indep_ds"))
    original_replace = os.replace

    def failing_replace(src, dst, *args, **kwargs):
        # 1. Fail during commit of tmp metadata
        if ".tmp_" in str(src) and (str(dst) == meta_target or str(dst).endswith("indep_ds.meta.json")):
            raise OSError("Injected failure during metadata commit")
        # 2. Fail during rollback of bak metadata
        if ".bak_" in str(src) and (str(dst) == meta_target or str(dst).endswith("indep_ds.meta.json")):
            raise OSError("Injected failure during metadata rollback")
        return original_replace(src, dst, *args, **kwargs)

    monkeypatch.setattr(os, "replace", failing_replace)

    with pytest.raises(OSError, match="Injected failure during metadata commit"):
        engine.store("indep_ds", v2_df, v2_meta)

    # Even though metadata rollback failed, Parquet rollback must have executed independently
    restored_df = engine.get("indep_ds")
    assert restored_df is not None
    pd.testing.assert_frame_equal(restored_df, v1_df)

    # The metadata backup file must NOT be deleted by cleanup, preserving recovery data
    meta_backups = [p for p in tmp_path.iterdir() if "indep_ds" in p.name and ".bak_" in p.name and "meta" in p.name]
    assert len(meta_backups) == 1, f"Expected metadata backup to be preserved, found: {meta_backups}"


def test_rollback_failure_preserves_surviving_backups_without_deletion(tmp_path, monkeypatch, caplog):
    """Verify that when rollback fails, surviving backup files are never deleted and CRITICAL log is emitted."""
    import os
    import logging
    from dataset_manager import DiskStorageEngine
    engine = DiskStorageEngine(storage_dir=str(tmp_path))
    v1_df = pd.DataFrame({"id": [1, 2], "val": ["V1_A", "V1_B"]})
    v1_meta = {"dataset_id": "surv_ds", "version": 1}

    engine.store("surv_ds", v1_df, v1_meta)

    v2_df = pd.DataFrame({"id": [10, 20], "val": ["V2_X", "V2_Y"]})
    v2_meta = {"dataset_id": "surv_ds", "version": 2}

    df_target = str(engine._df_path("surv_ds"))
    meta_target = str(engine._meta_path("surv_ds"))
    original_replace = os.replace

    def failing_replace(src, dst, *args, **kwargs):
        # Fail on metadata commit
        if ".tmp_" in str(src) and (str(dst) == meta_target or str(dst).endswith("surv_ds.meta.json")):
            raise OSError("Injected failure during metadata commit")
        # Fail on Parquet rollback
        if ".bak_" in str(src) and (str(dst) == df_target or str(dst).endswith("surv_ds.parquet")):
            raise OSError("Injected failure during Parquet rollback")
        return original_replace(src, dst, *args, **kwargs)

    monkeypatch.setattr(os, "replace", failing_replace)

    with caplog.at_level(logging.CRITICAL):
        with pytest.raises(OSError, match="Injected failure during metadata commit"):
            engine.store("surv_ds", v2_df, v2_meta)

    # Parquet backup MUST NOT be unlinked by cleanup
    parquet_backups = [p for p in tmp_path.iterdir() if "surv_ds" in p.name and ".bak_" in p.name and "parquet" in p.name]
    assert len(parquet_backups) == 1, f"Expected Parquet backup to be preserved, found: {parquet_backups}"
    # Verify critical log was recorded
    critical_logs = [r for r in caplog.records if r.levelno >= logging.CRITICAL]
    assert any("Failed to restore Parquet backup" in r.message for r in critical_logs)


def test_failure_during_metadata_backup_restores_parquet_backup(tmp_path, monkeypatch):
    """Verify that failure during metadata backup cleanly restores any Parquet backup already created."""
    import os
    from dataset_manager import DiskStorageEngine
    engine = DiskStorageEngine(storage_dir=str(tmp_path))
    v1_df = pd.DataFrame({"id": [1, 2], "val": ["V1_A", "V1_B"]})
    v1_meta = {"dataset_id": "pre_meta_ds", "version": 1}

    engine.store("pre_meta_ds", v1_df, v1_meta)

    v2_df = pd.DataFrame({"id": [10, 20], "val": ["V2_X", "V2_Y"]})
    v2_meta = {"dataset_id": "pre_meta_ds", "version": 2}

    original_replace = os.replace

    def fail_on_backup_meta(src, dst, *args, **kwargs):
        # Trigger failure when moving original meta_path to bak_meta_path
        if ".bak_" in str(dst) and "meta.json" in str(dst):
            raise OSError("Injected failure creating metadata backup")
        return original_replace(src, dst, *args, **kwargs)

    monkeypatch.setattr(os, "replace", fail_on_backup_meta)

    with pytest.raises(OSError, match="Injected failure creating metadata backup"):
        engine.store("pre_meta_ds", v2_df, v2_meta)

    # Parquet file must be restored back to destination
    assert engine.exists("pre_meta_ds") is True
    pd.testing.assert_frame_equal(engine.get("pre_meta_ds"), v1_df)
    assert engine.get_metadata("pre_meta_ds")["version"] == 1


def test_list_ids_excludes_temporary_and_backup_files(tmp_path):
    """Verify that list_ids() excludes temporary, backup, and incomplete dataset files."""
    from dataset_manager import DiskStorageEngine
    engine = DiskStorageEngine(storage_dir=str(tmp_path))

    # 1. Valid dataset: both parquet and meta exist
    valid_df = pd.DataFrame({"id": [1]})
    valid_meta = {"dataset_id": "valid_ds"}
    engine.store("valid_ds", valid_df, valid_meta)

    # 2. Incomplete dataset: parquet exists, but no meta.json
    orphan_df_path = engine.storage_dir / "orphan_ds.parquet"
    valid_df.to_parquet(orphan_df_path, index=False)

    # 3. Stranded temporary files
    (engine.storage_dir / "valid_ds.parquet.tmp_12345").touch()
    (engine.storage_dir / "stranded.tmp_12345.parquet").touch()

    # 4. Stranded backup files
    (engine.storage_dir / "valid_ds.parquet.bak_12345").touch()
    (engine.storage_dir / "stranded.bak_12345.parquet").touch()

    # list_ids() must ONLY return valid datasets with both parquet and metadata
    active_ids = engine.list_ids()
    assert active_ids == ["valid_ds"], f"Expected only ['valid_ds'], found: {active_ids}"
